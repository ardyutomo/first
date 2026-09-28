"""Generate CV Novita Ilmaris in DOCX and PDF formats."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Image,
                                Table, TableStyle, HRFlowable)

HERE = os.path.dirname(os.path.abspath(__file__))
PHOTO = os.path.join(HERE, "foto.png")
NAVY = "#1F3A5F"
GOLD = "#B8912F"

NAME = "Novita Ilmaris, S.Kom., S.H., M.H."
TITLE = "Sekretaris Jenderal Kementerian Hak Asasi Manusia Republik Indonesia"

PROFIL = (
    "Aparatur Sipil Negara dengan pengalaman lebih dari dua dekade di bidang tata kelola "
    "pemerintahan, manajemen sumber daya manusia, keuangan, dan pengelolaan aset negara. "
    "Memadukan latar belakang teknologi informasi dan ilmu hukum untuk mendorong birokrasi "
    "yang modern, akuntabel, dan berorientasi pada pelayanan publik. Sejak 6 Januari 2025 "
    "dipercaya mengemban amanah sebagai Sekretaris Jenderal Kementerian Hak Asasi Manusia RI "
    "di bawah kepemimpinan Menteri Hak Asasi Manusia, Natalius Pigai."
)

PERAN = [
    "Mengoordinasikan perencanaan strategis, pengelolaan sumber daya, dan dukungan administrasi "
    "bagi seluruh unit kerja di lingkungan Kementerian Hak Asasi Manusia.",
    "Memastikan setiap program dan kebijakan dilaksanakan secara efektif, efisien, dan akuntabel.",
    "Menjembatani visi dan arahan Menteri dengan implementasi program di seluruh unit kerja, "
    "sehingga setiap kebijakan dan kegiatan selaras dengan prinsip-prinsip hak asasi manusia.",
    "Menjadi penghubung strategis antara pimpinan, unit kerja, dan masyarakat guna mewujudkan "
    "pelayanan publik yang prima.",
]

PENDIDIKAN = [
    ("2021", "Magister Hukum (M.H.)", "Universitas Islam Sultan Agung, Semarang"),
    ("2003", "Sarjana Hukum (S.H.)", "Universitas Indonesia"),
    ("1998", "Sarjana Komputer (S.Kom.)", "STMIK Budi Luhur"),
]

KARIR = [
    ("2025 – sekarang", "Sekretaris Jenderal, Kementerian Hak Asasi Manusia RI"),
    ("", "Sekretaris Direktorat Jenderal Hak Asasi Manusia"),
    ("", "Kepala Biro Pengelolaan Barang Milik Negara"),
    ("", "Kepala Divisi Administrasi, Kantor Wilayah Banten"),
    ("", "Kepala Divisi Administrasi, Kantor Wilayah Daerah Istimewa Yogyakarta"),
    ("", "Kepala Bagian Perencanaan dan Sistem Informasi, Biro Kepegawaian"),
    ("", "Kepala Bagian Tata Usaha Kepegawaian, Biro Kepegawaian"),
    ("", "Kepala Bagian Keuangan, Sekretariat BPSDM"),
    ("", "Kepala Bidang Penyelenggaraan"),
    ("1999", "Pegawai Negeri Sipil, Kementerian Hukum dan HAM RI"),
]

KARYA = (
    "Aktif menulis gagasan dan pemikiran di berbagai media cetak maupun daring. "
    "Karya terbaru berupa buku berjudul <i>Memimpin Sistem, Menghasilkan Dampak</i>."
)
KARYA_PLAIN = KARYA.replace("<i>", "").replace("</i>", "")

SECTIONS = ["Profil Profesional", "Peran dan Tanggung Jawab Strategis",
            "Riwayat Pendidikan", "Riwayat Jabatan", "Karya dan Publikasi"]


# ---------------- DOCX ----------------
def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(NAVY[1:])
    # bottom border
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", GOLD[1:])):
        b.set(qn("w:" + k), v)
    bdr.append(b)
    pPr.append(bdr)


def build_docx(path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = t.rows[0].cells
    c0.width, c1.width = Cm(4.5), Cm(12)
    c0.paragraphs[0].add_run().add_picture(PHOTO, width=Cm(4))
    p = c1.paragraphs[0]
    r = p.add_run(NAME)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(NAVY[1:])
    p2 = c1.add_paragraph()
    r2 = p2.add_run(TITLE)
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = RGBColor.from_string(GOLD[1:])
    r2.bold = True

    heading(doc, SECTIONS[0])
    p = doc.add_paragraph(PROFIL)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading(doc, SECTIONS[1])
    for item in PERAN:
        p = doc.add_paragraph(item, style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading(doc, SECTIONS[2])
    for yr, deg, inst in PENDIDIKAN:
        p = doc.add_paragraph()
        r = p.add_run(f"{yr}\t")
        r.bold = True
        r = p.add_run(deg)
        r.bold = True
        p.add_run(f" — {inst}")
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.5))
        p.paragraph_format.space_after = Pt(2)

    heading(doc, SECTIONS[3])
    for yr, job in KARIR:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.5))
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{yr}\t")
        r.bold = True
        p.add_run(job)

    heading(doc, SECTIONS[4])
    p = doc.add_paragraph()
    p.add_run("Aktif menulis gagasan dan pemikiran di berbagai media cetak maupun daring. "
              "Karya terbaru berupa buku berjudul ")
    p.add_run("Memimpin Sistem, Menghasilkan Dampak").italic = True
    p.add_run(".")
    doc.save(path)


# ---------------- PDF ----------------
def build_pdf(path):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title=f"CV - {NAME}", author=NAME)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=10.3, leading=14.5,
                          alignment=TA_JUSTIFY)
    h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=12, leading=15,
                       textColor=colors.HexColor(NAVY), spaceBefore=12, spaceAfter=2)
    name = ParagraphStyle("n", fontName="Helvetica-Bold", fontSize=19, leading=23,
                          textColor=colors.HexColor(NAVY))
    sub = ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=11.5, leading=15,
                         textColor=colors.HexColor(GOLD), spaceBefore=6)
    bullet = ParagraphStyle("bl", parent=body, leftIndent=14, bulletIndent=2, spaceAfter=3)

    def sec(title):
        return [Paragraph(title.upper(), h),
                HRFlowable(width="100%", thickness=1.2, color=colors.HexColor(GOLD),
                           spaceAfter=6)]

    story = []
    header = Table([[Image(PHOTO, 4 * cm, 4 * cm),
                     [Paragraph(NAME, name), Paragraph(TITLE, sub)]]],
                   colWidths=[4.6 * cm, None])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("LEFTPADDING", (0, 0), (0, 0), 0)]))
    story += [header, Spacer(1, 4)]
    story += sec(SECTIONS[0]) + [Paragraph(PROFIL, body)]
    story += sec(SECTIONS[1]) + [Paragraph(x, bullet, bulletText="•") for x in PERAN]

    def two_col(rows):
        t = Table(rows, colWidths=[3.5 * cm, None])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                               ("TOPPADDING", (0, 0), (-1, -1), 1)]))
        return t

    left = ParagraphStyle("l", parent=body, fontName="Helvetica-Bold", alignment=TA_LEFT)
    lb = ParagraphStyle("lb", parent=body, alignment=TA_LEFT)
    story += sec(SECTIONS[2]) + [two_col(
        [[Paragraph(y, left), Paragraph(f"<b>{d}</b> — {i}", lb)] for y, d, i in PENDIDIKAN])]
    story += sec(SECTIONS[3]) + [two_col(
        [[Paragraph(y, left), Paragraph(j, lb)] for y, j in KARIR])]
    story += sec(SECTIONS[4]) + [Paragraph(KARYA, body)]
    doc.build(story)


if __name__ == "__main__":
    build_docx(os.path.join(HERE, "CV_Novita_Ilmaris.docx"))
    build_pdf(os.path.join(HERE, "CV_Novita_Ilmaris.pdf"))
    print("done")
